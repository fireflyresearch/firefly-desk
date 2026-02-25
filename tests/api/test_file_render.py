# Copyright 2026 Firefly Software Solutions Inc
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0

"""Tests for the file render endpoint."""

from __future__ import annotations

import io
import os
from unittest.mock import patch

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from flydesk.files.extractor import ContentExtractor
from flydesk.files.repository import FileUploadRepository
from flydesk.files.storage import LocalFileStorage
from flydesk.models.base import Base

# ---------------------------------------------------------------------------
# Helpers — create minimal office documents in memory
# ---------------------------------------------------------------------------


def _make_docx(paragraphs: list[str]) -> bytes:
    """Create a minimal DOCX file with the given paragraph texts."""
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(header: list[str], rows: list[list]) -> bytes:
    """Create a minimal XLSX file with a header row and data rows."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(header)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
async def client(tmp_path):
    env = {
        "FLYDESK_DATABASE_URL": "sqlite+aiosqlite:///:memory:",
        "FLYDESK_DEV_MODE": "true",
        "FLYDESK_AGENT_NAME": "Ember",
        "FLYDESK_FILE_STORAGE_PATH": str(tmp_path / "uploads"),
    }
    with patch.dict(os.environ, env):
        from flydesk.api.files import get_content_extractor, get_file_repo, get_file_storage
        from flydesk.server import create_app

        app = create_app()

        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        async with engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        session_factory = async_sessionmaker(engine, expire_on_commit=False)

        file_repo = FileUploadRepository(session_factory)
        file_storage = LocalFileStorage(str(tmp_path / "uploads"))
        content_extractor = ContentExtractor()

        app.dependency_overrides[get_file_repo] = lambda: file_repo
        app.dependency_overrides[get_file_storage] = lambda: file_storage
        app.dependency_overrides[get_content_extractor] = lambda: content_extractor

        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://test") as ac:
            yield ac

        await engine.dispose()


async def _upload(client: AsyncClient, filename: str, content: bytes, content_type: str) -> str:
    """Upload a file via the existing endpoint and return its ``file_id``."""
    resp = await client.post(
        "/api/chat/upload",
        files={"file": (filename, content, content_type)},
    )
    assert resp.status_code == 201, resp.text
    return resp.json()["id"]


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestRenderFile:
    """Tests for GET /api/files/{file_id}/render."""

    async def test_render_docx_paragraphs(self, client):
        """DOCX renders to HTML with paragraph text."""
        docx_bytes = _make_docx(["Hello world", "Second paragraph"])
        file_id = await _upload(
            client,
            "test.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        resp = await client.get(f"/api/files/{file_id}/render")
        assert resp.status_code == 200
        body = resp.json()
        assert body["type"] == "docx"
        assert "<p>Hello world</p>" in body["html"]
        assert "<p>Second paragraph</p>" in body["html"]

    async def test_render_xlsx_sheets(self, client):
        """XLSX renders to JSON with sheets, columns, and rows."""
        xlsx_bytes = _make_xlsx(["Name", "Age"], [["Alice", 30], ["Bob", 25]])
        file_id = await _upload(
            client,
            "data.xlsx",
            xlsx_bytes,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

        resp = await client.get(f"/api/files/{file_id}/render")
        assert resp.status_code == 200
        body = resp.json()
        assert body["type"] == "xlsx"
        assert len(body["sheets"]) >= 1
        sheet = body["sheets"][0]
        assert sheet["title"] == "Sheet1"
        assert len(sheet["columns"]) == 2
        assert sheet["columns"][0]["label"] == "Name"
        assert sheet["columns"][1]["label"] == "Age"
        assert len(sheet["rows"]) == 2
        assert sheet["rows"][0]["col_0"] == "Alice"

    async def test_render_missing_file_returns_404(self, client):
        """Render endpoint returns 404 for a non-existent file."""
        resp = await client.get("/api/files/nonexistent-id/render")
        assert resp.status_code == 404

    async def test_render_text_file(self, client):
        """Plain text files render with type='text'."""
        file_id = await _upload(client, "readme.txt", b"Hello text", "text/plain")

        resp = await client.get(f"/api/files/{file_id}/render")
        assert resp.status_code == 200
        body = resp.json()
        assert body["type"] == "text"
        assert body["text"] == "Hello text"

    async def test_render_pdf_returns_download_url(self, client):
        """PDF files render with type='pdf' and a download URL."""
        file_id = await _upload(client, "doc.pdf", b"%PDF-1.4 fake", "application/pdf")

        resp = await client.get(f"/api/files/{file_id}/render")
        assert resp.status_code == 200
        body = resp.json()
        assert body["type"] == "pdf"
        assert body["download_url"] == f"/api/files/{file_id}/download"

    async def test_render_image_returns_download_url(self, client):
        """Image files render with type='image' and a download URL."""
        file_id = await _upload(
            client, "photo.png", b"\x89PNG\r\n\x1a\n", "image/png"
        )

        resp = await client.get(f"/api/files/{file_id}/render")
        assert resp.status_code == 200
        body = resp.json()
        assert body["type"] == "image"
        assert body["download_url"] == f"/api/files/{file_id}/download"
